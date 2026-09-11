# Файл: src/reports/budget_exporter.py
import os
import numpy as np
import openpyxl


class ExcelExporter:
    """Суперстабильный модуль экспорта бюджета в Excel без формул и проблемных вызовов сетки"""

    @staticmethod
    def export_emc_budget(output_dir: str, task_name: str, method_name: str,
                          freq_range_str: str, budget_data: list):
        """Создает базовую таблицу данных с чистыми цифровыми значениями RSS без сбоев XML"""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "EMC Budget"

        # 💡 ИСПРАВЛЕНИЕ: Самый простой и совместимый способ включения сетки во всех версиях openpyxl
        # Он никогда не вызывает ошибку AttributeError: 'list' object...
        ws.sheet_view.showGridLines = True

        # 1. Записываем информационную шапку отчета
        ws.append([f"Инструментальный бюджет неопределенности: {method_name}"])
        ws.append([f"Исследуемый диапазон частот: {freq_range_str}"])
        ws.append([])  # Пустая строка-разделитель

        # 2. Записываем названия колонок таблицы
        headers = [
            "Источник неопределенности (Составляющая)",
            "Измерительное оборудование / Модель",
            "Исходное значение [дБ]",
            "Закон распределения (k)",
            "Стандартная неопределенность u(xi) [дБ]"
        ]
        ws.append(headers)

        # 3. Последовательно заполняем строки данными измерительных приборов
        start_data_row = 5
        sum_squares = 0.0  # Сюда будем накапливать квадраты для расчета RSS на Python

        for item in budget_data:
            u_std_value = float(item['u_std'])
            sum_squares += u_std_value ** 2  # Накапливаем сумму квадратов составляющих

            ws.append([
                item['name'],
                item['device'],
                item['raw_val'],
                item['dist'],
                u_std_value
            ])

        end_data_row = ws.max_row

        # 4. МЕТРОЛОГИЧЕСКИЙ РАСЧЕТ СИЛАМИ PYTHON (Абсолютная стабильность)
        u_combined_standard = np.sqrt(sum_squares)  # Суммарная стандартная неопределенность (k=1)
        u_expanded_lab = u_combined_standard * 2.0  # Расширенная неопределенность (k=2)

        comb_row = end_data_row + 1
        exp_row = end_data_row + 2

        # Записываем готовые, рассчитанные числа напрямую в ячейки Excel
        ws.cell(row=comb_row, column=1, value="СУММАРНАЯ СТАНДАРТНАЯ НЕОПРЕДЕЛЕННОСТЬ")
        ws.cell(row=comb_row, column=4, value="Расчет по методу RSS (k=1)")
        ws.cell(row=comb_row, column=5, value=round(u_combined_standard, 4))

        ws.cell(row=exp_row, column=1, value="РАСШИРЕННАЯ НЕОПРЕДЕЛЕННОСТЬ ЛАБОРАТОРИИ (U_lab)")
        ws.cell(row=exp_row, column=4, value="NORMAL (k = 2, 95%)")
        ws.cell(row=exp_row, column=5, value=round(u_expanded_lab, 4))

        # 5. Жесткое задание ширины столбцов под русский текст, чтобы ничего не обрезалось
        ws.column_dimensions['A'].width = 55
        ws.column_dimensions['B'].width = 45
        ws.column_dimensions['C'].width = 25
        ws.column_dimensions['D'].width = 30
        ws.column_dimensions['E'].width = 40

        # Сохранение итогового файла книги
        os.makedirs(output_dir, exist_ok=True)
        file_path = os.path.join(output_dir, f"Uncertainty_Budget_{task_name}.xlsx")
        wb.save(file_path)
        print(f"[ExcelExporter] Базовый цифровой отчет успешно сохранен: {file_path}")
